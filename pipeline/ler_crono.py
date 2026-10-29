"""Espelho Python de pbi/m/stgCrono.m: lê a planilha da cronograma de versão mais recente e confronta, obra × campo, com o
cadastro antigo (dados/cadastro/_antes-crono/obras_cadastro.csv, copiado à mão da cronograma de abr/26 antes da automação).

    python -X utf8 pipeline/ler_crono.py                              # pasta padrão da cronograma (dados/cronograma do repositório)
    python -X utf8 pipeline/ler_crono.py --pasta <pasta> [--antes <csv>]

Só leitura (openpyxl, read_only, data_only): a planilha da cronograma é do cliente e nunca é gravada.
Regras (mesmas de stgCrono.m; a descrição da consulta traz a evidência):
  Orçamento = ORÇ OBRA RESULTADO · M2 = M2 EQUIVALENTE (NBR) · Tipo = TIPO DE ORÇAMENTO com a inicial maiúscula
  Situação: realizado acumulado ÷ orçamento ≥ 99% Concluída · < 0,5% A iniciar · senão Em andamento
  Início: A iniciar → 1.º mês com previsto depois do mês de corte; demais → 1.º mês da aba FLUXO-FINACEIRO da obra
  Término: último mês com previsto diferente de zero na PREVXREA
Códigos de saída: 0 tudo igual · 1 com diferença (explicada no relatório) · 2 bloqueio (pasta ou arquivo sem leitura).
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[1]
PASTA_PADRAO = RAIZ / "dados" / "cronograma"
PREFIXO = "planilha de fluxo financeiro"
CADASTRO_ANTES = RAIZ / "dados" / "cadastro" / "_antes-crono" / "obras_cadastro.csv"
CADASTRO_ATUAL = RAIZ / "dados" / "cadastro" / "obras_cadastro.csv"   # tem nome_crono (vínculo com a cronograma)
MESES = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6, "jul": 7, "ago": 8, "set": 9, "out": 10,
         "nov": 11, "dez": 12}
ESTAGIOS = {"PRE", "EX", "VI", "LE"}
PREFIXOS_ABA = {"FLUXO-FINACEIRO", "FLUXO-FINANCEIRO"}


def norm(v) -> str:
    """fxNormaliza: minúsculas, sem acento, sem espaço duplicado."""
    if v is None:
        return ""
    s = unicodedata.normalize("NFKD", str(v).lower())
    s = "".join(c for c in s if not unicodedata.combining(c)).replace("º", "o").replace("ª", "a").replace("°", "")
    return " ".join(s.split())


def num(v):
    if v is None or isinstance(v, bool):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# Como o Excel.Workbook do Power Query entrega uma data: só a célula com formato de data EMBUTIDO do Excel (numFmtId 14 a
# 22: m/d/yyyy, d-mmm-yy, mmm-yy, …) chega como data; com formato personalizado ("mmm/yy", que o LibreOffice grava para
# todo formato), chega o número de série. Por isso o M (stgCrono, fxLerAba) aceita também um número de série plausível,
# e este espelho entrega o número no mesmo caso: assim ele pega essa classe de erro.
DATA_EMBUTIDA = range(14, 23)
ORIGEM_SERIAL = dt.datetime(1899, 12, 30)
SERIAL_MIN, SERIAL_MAX = 30000, 60000      # de 1982 a 2064: número de série plausível para um mês de obra


def valor_pq(c):
    """O valor da célula (openpyxl) como o Power Query o entrega: data com formato personalizado vira número de série."""
    v = getattr(c, "value", None)
    if isinstance(v, (dt.datetime, dt.date)):
        estilo = getattr(c, "style_array", None) or getattr(c, "_style", None)
        if estilo is None or estilo.numFmtId not in DATA_EMBUTIDA:
            d = v if isinstance(v, dt.datetime) else dt.datetime(v.year, v.month, v.day)
            return (d - ORIGEM_SERIAL).total_seconds() / 86400
    return v


def de_serial(v: float) -> dt.date:
    """Date.From(número) do M: 1899-12-30 + n dias."""
    return (ORIGEM_SERIAL + dt.timedelta(days=float(v))).date()


def eh_data(v) -> bool:
    """EhData do stgCrono: data, ou número de série plausível (a data que chegou como número)."""
    if isinstance(v, (dt.datetime, dt.date)):
        return True
    return isinstance(v, (int, float)) and not isinstance(v, bool) and SERIAL_MIN <= v <= SERIAL_MAX


def mes_de(v):
    if not eh_data(v):
        return None
    d = de_serial(v) if isinstance(v, (int, float)) else v
    return dt.date(d.year, d.month, 1)


def mes_da_celula(v):
    """Coluna B da aba FLUXO: data, ou "ACUMULADO-MM/AA" (conta como MM/AA)."""
    if eh_data(v):
        return mes_de(v)
    if isinstance(v, str) and v.strip().upper().startswith("ACUMULADO-"):
        partes = v.strip().split("-", 1)[1].split("/")
        try:
            mes, ano = int(partes[0]), int(partes[1])
            if 1 <= mes <= 12:
                return dt.date(2000 + ano, mes, 1)
        except (ValueError, IndexError):
            return None
    return None


def chave_versao(nome: str) -> tuple[int, int]:
    m = re.search(r"ATUALIZADO EM (\d{6})\.", nome.upper())
    if not m:
        return 0, 0
    d = m.group(1)
    return int(d[2:] + d[:2]), 1


def versao_txt(nome: str, modificado: dt.datetime) -> str:
    m = re.search(r"ATUALIZADO EM (\d{6})\.", nome.upper())
    if m:
        return f"{m.group(1)[:2]}/{m.group(1)[2:]}"
    return modificado.strftime("%m/%Y")


def chave_mes_aba(nome: str) -> int:
    tokens = [t for t in norm(nome).upper().split(" ") if t]
    mes = next((MESES[t.lower()] for t in tokens if t.lower() in MESES), None)
    ano = next((int(t) + (2000 if len(t) == 2 else 0) for t in tokens
                if t.isdigit() and len(t) in (2, 4)), None)
    return ano * 100 + mes if mes and ano else 0


def obra_da_aba(nome: str) -> str:
    tokens = [t for t in norm(nome).upper().split(" ") if t]
    if tokens and tokens[0] in PREFIXOS_ABA:
        tokens = tokens[1:]
    while tokens and tokens[0] in ESTAGIOS:
        tokens = tokens[1:]
    return " ".join(tokens)


def achar_arquivo(pasta: Path):
    if not pasta.is_dir():
        raise SystemExit(f"[BLOQUEIO] pasta da cronograma não abre: {pasta}")
    cand = [p for p in pasta.iterdir() if p.is_file() and norm(p.name).startswith(PREFIXO)
            and p.suffix.lower() == ".xlsx" and not p.name.startswith("~$")]
    if not cand:
        raise SystemExit(f"[BLOQUEIO] nenhuma planilha '{PREFIXO}' (.xlsx) em {pasta}")
    return sorted(cand, key=lambda p: (chave_versao(p.name)[0], p.stat().st_mtime), reverse=True)[0]


def linhas_usadas(ws) -> list[tuple]:
    """As linhas da aba a partir da primeira célula usada, como o Excel.Workbook do Power Query as entrega (uma coluna A
    vazia some: por isso as colunas são achadas pelo rótulo, nunca pela posição; data com formato personalizado chega
    como número de série, valor_pq)."""
    return [tuple(valor_pq(c) for c in linha) for linha in ws.iter_rows(min_row=ws.min_row, min_col=ws.min_column)]


def ler_fluxo(ws) -> list[dict] | None:
    """Curva da aba FLUXO-FINACEIRO: [{mes, custo}] pelas colunas MÊS e CUSTO PREVISTO (o cabeçalho numa das 15
    primeiras linhas usadas, em qualquer coluna); None sem o cabeçalho."""
    linhas = linhas_usadas(ws)
    h = next((i for i in range(min(15, len(linhas)))
              if {"mes", "custo previsto"} <= {norm(c) for c in linhas[i]}), None)
    if h is None:
        return None
    cab = [norm(c) for c in linhas[h]]
    cm, cc = cab.index("mes"), cab.index("custo previsto")
    out = []
    for r in linhas[h + 1:]:
        mes = mes_da_celula(r[cm] if len(r) > cm else None)
        if mes is not None:
            out.append({"mes": mes, "custo": num(r[cc]) if len(r) > cc else None})
    return out


def ler_crono(arquivo: Path) -> tuple[list[dict], list[dict]]:
    """Devolve (obras, avisos) como stgCrono.m: obras com Situacao, Inicio, Termino e campos da PREVXREA."""
    wb = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    versao = versao_txt(arquivo.name, dt.datetime.fromtimestamp(arquivo.stat().st_mtime))
    avisos: list[dict] = []
    previs = sorted([n for n in wb.sheetnames if norm(n).startswith("prevxrea")],
                    key=lambda n: (chave_mes_aba(n), n), reverse=True)
    if not previs:
        raise SystemExit(f"[BLOQUEIO] nenhuma aba PREVXREA em {arquivo.name}")
    nome_prev = previs[0]
    if len(previs) > 1:
        avisos.append({"obra": "", "aba": nome_prev, "regra": "aba_previxrea_varias", "sev": "info",
                       "det": f"{len(previs)} abas PREVXREA: vale a do mês mais recente no nome"})
    linhas = linhas_usadas(wb[nome_prev])
    h = next((i for i in range(min(15, len(linhas))) if norm(linhas[i][0]) == "obra"), None)
    if h is None:
        raise SystemExit(f"[BLOQUEIO] cabeçalho sem OBRA em {nome_prev}")
    cab, datas = linhas[h], linhas[h + 1]

    def col(criterio):
        return next((j for j, c in enumerate(cab) if criterio(norm(c))), None)

    c_obra = col(lambda r: r == "obra")
    c_tipo = col(lambda r: r.startswith("tipo de orcamento"))
    c_m2 = col(lambda r: r.startswith("m2 equivalente"))
    c_orc = col(lambda r: r.startswith("orc obra resultado"))
    c_real = col(lambda r: r == "realizado acumulado")
    c_prev = col(lambda r: r == "previsto")
    c_desv = col(lambda r: r == "desvio")
    faltam = [n for n, c in [("OBRA", c_obra), ("ORÇ OBRA RESULTADO", c_orc), ("PREVISTO", c_prev),
                             ("DESVIO", c_desv), ("REALIZADO ACUMULADO", c_real)] if c is None]
    if faltam:
        raise SystemExit(f"[BLOQUEIO] cabeçalho sem o rótulo {', '.join(faltam)}")
    cols_prev = [j for j in range(len(cab)) if j == c_prev or (j > c_desv and eh_data(datas[j]))]
    corte = mes_de(datas[c_prev])

    # abas FLUXO-FINACEIRO: primeiro mês da aba (coluna MÊS) por obra; soma do custo previsto (CUSTO PREVISTO)
    tabs: dict[str, list[dict]] = {}
    for n in sorted(wb.sheetnames):
        if norm(n).startswith("fluxo-fina"):
            meses = ler_fluxo(wb[n])
            if meses is None:
                avisos.append({"obra": obra_da_aba(n), "aba": n, "regra": "aba_fluxo_ilegivel", "sev": "alta",
                               "det": "cabeçalho com MÊS e CUSTO PREVISTO não encontrado"})
                meses = []
            tabs.setdefault(obra_da_aba(n), []).append({"aba": n, "meses": meses})

    obras, vistos = [], set()
    for r in linhas[h + 2:]:
        if r[c_obra] is None or str(r[c_obra]).strip() == "":
            break
        nome = str(r[c_obra]).strip().upper()
        if nome in vistos:
            avisos.append({"obra": nome, "aba": nome_prev, "regra": "obra_duplicada", "sev": "média",
                           "det": "aparece mais de uma vez na PREVXREA: vale a primeira linha"})
            continue
        vistos.add(nome)
        orc = num(r[c_orc])
        real = num(r[c_real]) or 0.0
        meses = []
        for j in cols_prev:
            v = num(r[j])
            if mes_de(datas[j]) and v not in (None, 0.0):
                meses.append(mes_de(datas[j]))
        pct = real / orc if orc else None
        situacao = None if pct is None else ("Concluída" if pct >= 0.99 else
                                             "A iniciar" if pct < 0.005 else "Em andamento")
        termino = max(meses) if meses else None
        ini_pre = min([m for m in meses if corte and m > corte], default=None)
        lista = tabs.get(nome, [])
        tab = lista[0] if lista else None
        ini_fluxo = min([x["mes"] for x in tab["meses"]], default=None) if tab else None
        inicio = ini_pre if situacao == "A iniciar" else ini_fluxo
        if len(lista) > 1:
            avisos.append({"obra": nome, "aba": lista[0]["aba"], "regra": "aba_fluxo_duplicada", "sev": "média",
                           "det": "mais de uma aba FLUXO-FINACEIRO para a obra: vale a primeira por nome"})
        if tab:
            soma = sum((x["custo"] or 0.0) for x in tab["meses"])
            if orc and orc > 0 and abs(soma / orc - 1) > 0.01:
                avisos.append({"obra": nome, "aba": tab["aba"], "regra": "curva_nao_fecha", "sev": "média",
                               "det": f"a curva soma {soma:,.2f} ({soma / orc * 100:.1f}% do orçamento)"})
        if tab is not None and not tab["meses"]:
            avisos.append({"obra": nome, "aba": tab["aba"], "regra": "aba_fluxo_sem_meses", "sev": "média",
                           "det": "a aba não tem mês lido na coluna B (datas): o início dessa obra fica vazio"})
        if situacao == "Em andamento" and tab is None:
            avisos.append({"obra": nome, "aba": nome_prev, "regra": "sem_aba_fluxo", "sev": "média",
                           "det": "obra em andamento sem aba FLUXO-FINACEIRO na planilha: o início fica vazio"})
        if situacao is None:
            avisos.append({"obra": nome, "aba": nome_prev, "regra": "situacao_indefinida", "sev": "média",
                           "det": "sem orçamento na PREVXREA"})
        obras.append({"obra": nome, "tipo": str(r[c_tipo]).strip().capitalize() if r[c_tipo] else None,
                      "orcamento": orc, "m2": num(r[c_m2]), "situacao": situacao, "inicio": inicio,
                      "termino": termino, "versao": versao, "arquivo": arquivo.name})
    return obras, avisos


def ym(d) -> str:
    return d.strftime("%Y-%m") if d else ""


def ler_csv(caminho: Path) -> dict[str, dict]:
    with caminho.open(encoding="utf-8-sig") as fh:
        return {r["obra"].strip().upper(): r for r in csv.DictReader(fh, delimiter=";")}


def main(argv=None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--pasta", type=Path, default=PASTA_PADRAO)
    ap.add_argument("--antes", type=Path, default=CADASTRO_ANTES)
    a = ap.parse_args(argv)
    try:
        arquivo = achar_arquivo(a.pasta)
        obras, avisos = ler_crono(arquivo)
    except SystemExit as e:
        print(e)
        return 2
    print(f"arquivo: {arquivo.name} · versão {obras[0]['versao'] if obras else '?'} · {len(obras)} obras na cronograma")
    antes = ler_csv(a.antes)
    nome_na_crono = {n: r.get("nome_crono") or n for n, r in ler_csv(CADASTRO_ATUAL).items()}
    campos = ["situacao", "tipo", "orcamento", "m2", "inicio", "termino"]
    crono = {o["obra"]: o for o in obras}
    cabec = ["obra", "situação", "tipo", "orçamento", "m²", "início", "término"]
    linhas_tab = []
    diferencas = 0
    for nome, velho in antes.items():
        g = crono.get(nome_na_crono.get(nome, nome))
        if g is None:
            linhas_tab.append([nome, "sem linha na cronograma"] + [""] * 5)
            diferencas += 1
            continue
        csv_v = {"situacao": velho["situacao"], "tipo": velho["tipo_orcamento_crono"],
                 "orcamento": round(float(velho["orcamento_crono_abr26"])),
                 "m2": round(float(velho["m2_nbr"])) if velho["m2_nbr"] else None,
                 "inicio": velho["inicio"] or "", "termino": velho["termino_previsto"] or ""}
        gt_v = {"situacao": g["situacao"] or "", "tipo": g["tipo"] or "",
                "orcamento": round(g["orcamento"]) if g["orcamento"] is not None else None,
                "m2": round(g["m2"]) if g["m2"] is not None else None,
                "inicio": ym(g["inicio"]), "termino": ym(g["termino"])}
        celulas = []
        for c in campos:
            cv = "" if csv_v[c] is None else csv_v[c]
            gv = "" if gt_v[c] is None else gt_v[c]
            if str(cv) == str(gv):
                celulas.append("igual")
            else:
                diferencas += 1
                celulas.append(f"{cv or '—'} → {gv or '—'}")
        linhas_tab.append([nome] + celulas)
    larguras = [max(len(str(x[i])) for x in [cabec] + linhas_tab) for i in range(len(cabec))]
    for x in [cabec] + linhas_tab:
        print(" | ".join(str(v).ljust(larguras[i]) for i, v in enumerate(x)))
    for o in obras:
        if o["obra"] not in set(nome_na_crono.values()):
            print(f"{o['obra']}: só na cronograma (fora do cadastro de antes)")
    print()
    print("avisos (como em fConferencia):")
    for av in avisos:
        print(f"  {av['obra'] or '(cronograma)':10} {av['regra']:22} {av['sev']:6} {av['det']}")
    print()
    print(f"diferenças: {diferencas} (ver a tabela; orçamento e m² comparados em inteiro, como estava no CSV)")
    return 1 if diferencas else 0


if __name__ == "__main__":
    sys.exit(main())
