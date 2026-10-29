"""Confere os dados fictícios (dados/) com o gabarito do gerador e mostra a prévia da Nota da obra.

    python -I -X utf8 pipeline/conferir_ficticios.py            # sai 1 se algo não bater

1. Análise Crítica: o extrator (pipeline/extrair_analise_critica.py, espelho de fxLerAba.m e fAnaliseCritica.m) lê as
   planilhas pelo cadastro, como o Power Query; cada campo tem de bater com dados/_gabarito/fato_analise_critica.csv
   (R$ ± 0,01; % exato), o resultado salvo pela planilha tem de bater com a regra (fora o resultado digitado de
   propósito) e os avisos têm de ser só os injetados pelo gerador.
2. Base complementar 2.0: leitura independente por rótulo (âncora "Versão do modelo" na coluna B, valor na D, como
   stgAbasADO.m) contra dados/_gabarito/base_complementar.csv.
3. Cronograma: pipeline/ler_crono.py (espelho de stgCrono.m) dá início, término e situação; a curva da aba FLUXO fecha
   no orçamento e bate com o previsto acumulado das abas da Análise Crítica.
4. Prévia da Nota da obra (régua de docs/08, faixas do dicionário NOTA de pbi/scripts/contas_base.py), mês a mês, de
   out/25 a set/26, gravada em dados/_gabarito/nota_previa.csv: o oráculo da F3. Confere o roteiro (quem fica em crítico
   e quando).
"""
from __future__ import annotations

import csv
import sys
from datetime import date
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[1]
DADOS = RAIZ / "dados"
GAB = DADOS / "_gabarito"
sys.path.insert(0, str(RAIZ / "pipeline"))
sys.path.insert(0, str(RAIZ / "pbi" / "scripts"))
import extrair_analise_critica as ex  # noqa: E402
import ler_crono as lc  # noqa: E402
from contas_base import NOTA, nota_pontos  # noqa: E402

HISTORICO = "2025-10-01"
CAMPOS_RS = ["orcamento", "credito", "saldo_contrato", "adiant_servicos", "adiant_materiais", "estoque", "equipamentos",
             "receita_personalizacao", "contas_pagar", "desembolsado_bruto", "itens_irregulares", "desembolsado_corrigido",
             "divida", "resultado", "resultado_recalculado", "previsto_mes_rs", "desembolsado_mes_rs"]
CAMPOS_PCT = ["prev_mes", "exec_mes", "prev_acum", "exec_acum"]
# avisos que o gerador injeta de propósito (item 9 da régua e fConferencia)
AVISOS_ESPERADOS = {("REDINHA", "2026-05-01", "conta_diverge"), ("JANGADA", "2026-03-01", "previsto_regrediu"),
                    ("JANGADA", "2026-04-01", "previsto_reprogramado"), ("JANGADA", "2026-04-01", "tipo_orcamento_mudou"),
                    ("TIROL", "2026-04-01", "mes_divergente")}
# roteiro (pedido de 28/10/2026): o status de cada obra em set/26 e em out/25
ROTEIRO = {("2026-09-01", "REDINHA"): "crítico", ("2026-09-01", "POTENGI"): "crítico",
           ("2026-09-01", "JANGADA"): "atenção", ("2026-09-01", "CAJUEIRO"): "estável",
           ("2026-09-01", "PONTA NEGRA"): "estável", ("2026-09-01", "DUNAS"): "estável",
           ("2026-09-01", "TIROL"): "estável", ("2025-10-01", "DUNAS"): "atenção",
           ("2025-10-01", "REDINHA"): "atenção", ("2025-10-01", "POTENGI"): "atenção"}
STATUS = {0: "estável", 1: "atenção", 2: "crítico"}

falhas: list[str] = []


def falha(msg: str):
    falhas.append(msg)


def ler_csv(p: Path) -> list[dict]:
    with p.open(encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh, delimiter=";"))


def f(v):
    return None if v in (None, "") else float(v)


# ---------------------------------------------------------------------------------------------------------------------
def conferir_ac() -> dict[tuple[str, str], dict]:
    arquivos, obras = ex.escolher_por_cadastro(DADOS / "analise-critica", DADOS / "cadastro" / "obras_cadastro.csv")
    linhas, itens, avisos = ex.processar(arquivos, obras)
    lidas = {(l["obra"], l["mes"]): l for l in linhas}
    gab = ler_csv(GAB / "fato_analise_critica.csv")
    n = 0
    for g in gab:
        l = lidas.get((g["obra"], g["mes"]))
        if l is None:
            falha(f"AC {g['obra']} {g['mes']}: aba não lida")
            continue
        for c in CAMPOS_RS:
            a, b = f(g[c]), l.get(c)
            if (a is None) != (b is None) or (a is not None and abs(a - b) > 0.011):
                falha(f"AC {g['obra']} {g['mes']} {c}: gabarito {a} × lido {b}")
        for c in CAMPOS_PCT:
            if abs(f(g[c]) - l[c]) > 1e-9:
                falha(f"AC {g['obra']} {g['mes']} {c}: gabarito {g[c]} × lido {l[c]}")
        if (g["tipo_orcamento"] or None) != l["tipo_orcamento"]:
            falha(f"AC {g['obra']} {g['mes']} tipo: {g['tipo_orcamento']} × {l['tipo_orcamento']}")
        if (g["justificativa"] or None) != l["justificativa"]:
            falha(f"AC {g['obra']} {g['mes']} justificativa: {g['justificativa']!r} × {l['justificativa']!r}")
        if l["resultado_excel"] is None:
            falha(f"AC {g['obra']} {g['mes']}: resultado sem valor salvo (o LibreOffice não recalculou?)")
        n += 1
    achados = {(a["obra"], a["mes"], a["regra"]) for a in avisos}
    for a in sorted(achados - AVISOS_ESPERADOS):
        falha(f"aviso inesperado: {a}")
    for a in sorted(AVISOS_ESPERADOS - achados):
        falha(f"aviso injetado e não achado: {a}")
    print(f"1. Análise Crítica: {n} abas de {len(arquivos)} planilhas conferidas · {len(itens)} linhas de tabelas · "
          f"{len(achados)} avisos (os {len(AVISOS_ESPERADOS)} injetados)")
    return lidas


# ---------------------------------------------------------------------------------------------------------------------
def conferir_bases() -> dict[tuple[str, str], dict]:
    rotulos = {ex.norm(r): c for c, r, _f in NOTA["base"]}
    lidas = {}
    for arq in sorted((DADOS / "bases-complementares").glob("*.xlsx")):
        obra = arq.stem.split(" - ", 1)[1].upper()
        wb = openpyxl.load_workbook(arq, data_only=True, read_only=True)
        for ws in wb.worksheets:
            mes = ex.mes_da_aba(ws.title)
            if mes is None:
                continue
            linhas = list(ws.iter_rows(values_only=True))
            ancora = next(((i, j) for i, ln in enumerate(linhas) for j, v in enumerate(ln)
                           if isinstance(v, str) and ex.norm(v) == "versao do modelo"), None)
            if ancora is None:
                falha(f"base {obra} {ws.title}: sem 'Versão do modelo'")
                continue
            col = ancora[1]
            reg = {"versao": linhas[ancora[0]][col + 2]}
            for ln in linhas:
                r = ln[col] if col < len(ln) else None
                if isinstance(r, str) and ex.norm(r) in rotulos:
                    reg[rotulos[ex.norm(r)]] = ln[col + 2]
            lidas[(obra, mes.isoformat())] = reg
    gab = ler_csv(GAB / "base_complementar.csv")
    for g in gab:
        l = lidas.get((g["obra"], g["mes"]))
        if l is None:
            falha(f"base {g['obra']} {g['mes']}: aba não lida")
            continue
        if l["versao"] != NOTA["versao_base"]:
            falha(f"base {g['obra']} {g['mes']}: versão {l['versao']}")
        for c, _r, (lo, hi) in NOTA["base"]:
            v = l.get(c)
            if v is None or abs(float(v) - float(g[c])) > 1e-9:
                falha(f"base {g['obra']} {g['mes']} {c}: gabarito {g[c]} × lido {v}")
            elif not lo <= float(v) <= hi:
                falha(f"base {g['obra']} {g['mes']} {c}: {v} fora da faixa válida")
        if float(l["AprovadosPrimeira"]) > float(l["ServicosVerificados"]) or \
                float(l["NR18Conformes"]) > float(l["NR18Aplicaveis"]) or float(l["PedidosNoPrazo"]) > float(l["PedidosPrevistos"]):
            falha(f"base {g['obra']} {g['mes']}: parte maior que o todo")
    print(f"2. Base complementar {NOTA['versao_base']}: {len(gab)} abas conferidas")
    return lidas


# ---------------------------------------------------------------------------------------------------------------------
def conferir_crono(ac: dict) -> dict[str, dict]:
    arq = lc.achar_arquivo(DADOS / "cronograma")
    obras, avisos = lc.ler_crono(arq)
    if avisos:
        falha(f"cronograma com avisos: {avisos}")
    cad = {c["obra"]: c for c in ler_csv(DADOS / "cadastro" / "obras_cadastro.csv")}
    wb = openpyxl.load_workbook(arq, data_only=True, read_only=True)
    curvas: dict[str, dict] = {}
    for o in obras:
        nome = o["obra"] if "obra" in o else o.get("ObraCrono")
        aba = next(n for n in wb.sheetnames if lc.obra_da_aba(n) == nome and lc.norm(n).startswith("fluxo-fina"))
        acum, curva = 0.0, {}
        # como o Power Query lê: a partir da primeira célula usada, MÊS e CUSTO PREVISTO pelo rótulo
        for ln in lc.ler_fluxo(wb[aba]) or falha(f"curva {nome}: aba {aba} sem MÊS e CUSTO PREVISTO") or []:
            acum += float(ln["custo"] or 0)
            curva[ln["mes"]] = acum
        orc = o["orcamento"]
        if abs(acum / orc - 1) > 0.001:
            falha(f"curva {nome}: soma {acum:,.2f} × orçamento {orc:,.2f}")
        curvas[nome] = {"inicio": o["inicio"], "termino": o["termino"], "orc": orc,
                        "p": {m: v / orc for m, v in curva.items()}, "situacao": o["situacao"]}
        if o["situacao"] != "Em andamento":
            falha(f"cronograma {nome}: situação {o['situacao']}")
        if nome not in cad:
            falha(f"cronograma {nome}: fora do cadastro")
    for (obra, mes), l in ac.items():
        c = curvas[obra]
        m = date.fromisoformat(mes)
        if m in c["p"] and l["prev_acum"] is not None and abs(round(c["p"][m], 4) - l["prev_acum"]) > 0.00015 and \
                (obra, mes) != ("JANGADA", "2026-03-01"):
            falha(f"previsto {obra} {mes}: curva {c['p'][m]:.4%} × aba {l['prev_acum']:.4%}")
    print(f"3. Cronograma: {len(obras)} obras · curvas fecham no orçamento e batem com o previsto das abas")
    return curvas


# ---------------------------------------------------------------------------------------------------------------------
def atraso_meses(curva: dict, exec_acum: float, mes: date) -> float:
    """Meses decorridos − meses agregados (Earned Schedule): ponto da curva prevista em que o previsto acumulado iguala
    o executado de hoje, interpolado entre os meses."""
    inicio = min(curva["p"])
    meses_c = sorted(curva["p"])
    pts = [(0, 0.0)] + [(i + 1, curva["p"][m]) for i, m in enumerate(meses_c)]
    t = (mes.year - inicio.year) * 12 + mes.month - inicio.month + 1
    es = pts[-1][0]
    for (k0, p0), (k1, p1) in zip(pts, pts[1:]):
        if p0 <= exec_acum <= p1:
            es = k0 + (exec_acum - p0) / (p1 - p0) if p1 > p0 else k0
            break
    return t - es


def previa_nota(ac: dict, bases: dict, curvas: dict):
    linhas, por = [], {}
    meses_h = sorted({m for (_o, m) in ac if m >= HISTORICO})
    ant_por_obra: dict[str, dict] = {}
    for (obra, mes) in sorted(ac, key=lambda k: (k[0], k[1])):
        l = ac[(obra, mes)]
        ant = ant_por_obra.get(obra)
        ant_por_obra[obra] = l
        if mes < HISTORICO:
            continue
        m = date.fromisoformat(mes)
        b = bases[(obra, mes)]
        ca = (l["desembolsado_corrigido"] + l["contas_pagar"] + l["divida"] - l["estoque"] - l["adiant_materiais"]
              - l["adiant_servicos"] - l["receita_personalizacao"])
        idc = l["credito"] / ca
        atraso = atraso_meses(curvas[obra], l["exec_acum"], m)
        desvio = (l["desembolsado_mes_rs"] / l["previsto_mes_rs"] - 1) if l["previsto_mes_rs"] else None
        caixa = abs(desvio) if desvio is not None else None
        janela = [bases.get((obra, date(m.year - (1 if m.month - k <= 0 else 0), (m.month - k - 1) % 12 + 1, 1).isoformat()))
                  for k in range(NOTA["janela_acidentes"])]
        janela = [j for j in janela if j]
        tf = 1e6 * sum(float(j["Acidentes"]) for j in janela) / sum(float(j["HHT"]) for j in janela)
        checagens = (int(abs(l["resultado"] - l["resultado_recalculado"]) < 1)
                     + int(caixa is None or caixa <= 0.10 or bool(l["justificativa"]))
                     + int(ant is None or ant["prev_acum"] is None or l["prev_acum"] >= ant["prev_acum"] - 1e-9))
        x = {"idc": idc, "atraso": atraso, "ppc": float(b["PPC"]), "caixa": caixa,
             "inspecao": float(b["AprovadosPrimeira"]) / float(b["ServicosVerificados"]), "acidentes": tf,
             "nr18": float(b["NR18Conformes"]) / float(b["NR18Aplicaveis"]),
             "entregas": float(b["PedidosNoPrazo"]) / float(b["PedidosPrevistos"]), "analise": checagens}
        pts = {k: nota_pontos(NOTA[k], v) for k, v in x.items()}
        nota = sum(pts.values())
        st = nota_pontos(NOTA["status"], nota)
        trava = idc < NOTA["trava"]["idc_abaixo"] or atraso > NOTA["trava"]["atraso_acima"]
        if trava:
            st = 2
        por[(mes, obra)] = (nota, STATUS[st], trava)
        linhas.append({"obra": obra, "mes": mes, **{f"x_{k}": round(v, 6) if v is not None else None for k, v in x.items()},
                       **{f"pts_{k}": v for k, v in pts.items()}, "nota": round(nota, 2), "trava": int(trava),
                       "status": STATUS[st]})
    with (GAB / "nota_previa.csv").open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(linhas[0].keys()), delimiter=";")
        w.writeheader()
        w.writerows(linhas)
    obras = sorted({o for (_m, o) in por})
    print("4. Prévia da Nota da obra (régua de docs/08), nota e status por mês (* = trava):")
    print("   " + "obra".ljust(12) + "".join(m[2:7].rjust(8) for m in meses_h))
    for o in obras:
        cel = []
        for m in meses_h:
            if (m, o) not in por:
                cel.append("".rjust(8))
                continue
            nota, st, trava = por[(m, o)]
            cel.append(f"{nota:5.2f}{'*' if trava else ' '}{st[0].upper()}".rjust(8))
        print("   " + o.ljust(12) + "".join(cel))
    for (m, o), esperado in ROTEIRO.items():
        if por.get((m, o), (None, None))[1] != esperado:
            falha(f"roteiro: {o} em {m[:7]} devia estar {esperado}, está {por.get((m, o), (None, '—'))[1]}")


def main() -> int:
    ac = conferir_ac()
    bases = conferir_bases()
    curvas = conferir_crono(ac)
    if not falhas:      # sem a cronograma lida (curva vazia, datas que não viram data), a prévia da nota não tem prazo
        previa_nota(ac, bases, curvas)
    if falhas:
        print(f"\n{len(falhas)} diferença(s):")
        for x in falhas[:60]:
            print("  -", x)
        return 1
    print("\nok: tudo bate com o gabarito e com o roteiro")
    return 0


if __name__ == "__main__":
    sys.exit(main())
