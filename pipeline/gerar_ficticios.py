"""Gera os dados fictícios do portfólio: uma construtora inventada de Natal/RN com sete empreendimentos. Tudo o que sai
daqui é inventado (obras, clientes, títulos, valores); os bairros são bairros reais de Natal.

    python -X utf8 pipeline/gerar_ficticios.py                 # gera em dados/ e recalcula as fórmulas (LibreOffice)
    python -X utf8 pipeline/gerar_ficticios.py --sem-recalculo # só o openpyxl (fórmulas sem valor salvo)

Saídas em dados/ (o layout é o que o Power Query lê; docs/02 e docs/08 §7):
    cadastro/obras_cadastro.csv                         cadastro das obras (CSV ";" UTF-8): bairro, unidades e a
                                                        data de entrega contratual entram nesta versão
    analise-critica/<OBRA>/Analise Critica <OBRA>.xlsx  uma aba por mês (SET25 … SET26), lida por rótulo (fxLerAba.m)
    bases-complementares/Base complementar - <OBRA>.xlsx  base 2.0, uma aba por mês (OUT 2024 … SET 2026; a taxa de
                                                        frequência de acidentes olha 12 meses para trás)
    cronograma/PLANILHA DE FLUXO FINANCEIRO ATUALIZADO EM 092026.xlsx   PREVXREA de set/26 e uma aba FLUXO-FINACEIRO
                                                        por obra com a curva prevista mês a mês (linha de base)
    _gabarito/*.csv                                     o que o gerador quis escrever, para pipeline/conferir_ficticios.py

As planilhas saem com fórmulas, como a planilha da obra. O Power Query lê o valor salvo de cada fórmula, e quem calcula e
salva esse valor é o LibreOffice em modo headless (no lugar do Excel do gerador antigo).

Roteiro (docs/08 e pedido do Alisson, 28/10/2026): out/25 a set/26 fechados; obras em fases diferentes (TIROL começando
em fev/26, CAJUEIRO terminando); REDINHA e POTENGI entram em crítico (custo e prazo); DUNAS melhora; a quadra chuvosa de
Natal (mar a jul, pico em abril) marca o físico; a dívida da obra desconta do resultado (planilha versão 01).
"""
from __future__ import annotations

import argparse
import csv
import html
import random
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

RAIZ = Path(__file__).resolve().parents[1]
DADOS = RAIZ / "dados"
sys.path.insert(0, str(RAIZ / "pbi" / "scripts"))
from contas_base import NOTA  # noqa: E402  (campos da base complementar 2.0: fonte única)

CORTE = date(2026, 9, 1)         # último mês fechado
INICIO_AC = date(2025, 9, 1)     # primeira aba da Análise Crítica (o "mês anterior" de out/25)
INICIO_BASE = date(2024, 10, 1)  # primeira aba das bases (12 meses antes de out/25, para a taxa de frequência)
INCC_MES = 0.0042                # correção do orçamento projetado por mês (INCC no a realizar)
SIGLAS = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]
# dias impraticáveis por chuva em Natal (quadra chuvosa de mar a jul, pico em abril; normal do INMET para abril 240,5 mm)
CHUVA = {1: 2, 2: 3, 3: 5, 4: 9, 5: 8, 6: 7, 7: 5, 8: 2, 9: 1, 10: 0, 11: 0, 12: 1}
# quanto da chuva vira atraso em cada fase (fração do mês perdida por dia impraticável ÷ 21 dias úteis)
SENS_CHUVA = {"mobilizacao": 0.9, "fundacao": 0.9, "estrutura": 0.6, "alvenaria": 0.35, "acabamento": 0.15,
              "entrega": 0.05}
EFETIVO_POR_MI = {"mobilizacao": 1.3, "fundacao": 2.1, "estrutura": 3.1, "alvenaria": 3.7, "acabamento": 3.3,
                  "entrega": 1.6}
ESTOQUE_FRAC = {"mobilizacao": 0.002, "fundacao": 0.005, "estrutura": 0.010, "alvenaria": 0.012, "acabamento": 0.009,
                "entrega": 0.002}

random.seed(20261028)


# ---------------------------------------------------------------------------------------------------------------------
# Datas e curvas
# ---------------------------------------------------------------------------------------------------------------------
def add_mes(d: date, n: int) -> date:
    y, m = divmod(d.month - 1 + n, 12)
    return date(d.year + y, m + 1, 1)


def meses_entre(a: date, b: date) -> int:
    return (b.year - a.year) * 12 + b.month - a.month


def meses(a: date, b: date) -> list[date]:
    return [add_mes(a, i) for i in range(meses_entre(a, b) + 1)]


def chave(d: date) -> str:
    return f"{d.year}-{d.month:02d}"


def interp(nos: dict[str, float], d: date) -> float:
    """Valor no mês d, interpolado entre os nós {"AAAA-MM": valor}; fora deles, o nó da ponta."""
    pts = sorted((date(int(k[:4]), int(k[5:]), 1), v) for k, v in nos.items())
    if d <= pts[0][0]:
        return pts[0][1]
    for (d0, v0), (d1, v1) in zip(pts, pts[1:]):
        if d0 <= d <= d1:
            return v0 + (v1 - v0) * meses_entre(d0, d) / meses_entre(d0, d1)
    return pts[-1][1]


def curva_s(x: float) -> float:
    """Curva S do previsto acumulado (fração do orçamento) na fração x do prazo."""
    x = min(max(x, 0.0), 1.0)
    return x * x * (3 - 2 * x)


def fase(e: float) -> str:
    return ("mobilizacao" if e < 0.06 else "fundacao" if e < 0.20 else "estrutura" if e < 0.50
            else "alvenaria" if e < 0.80 else "acabamento" if e < 0.97 else "entrega")


def nome_aba_ac(d: date) -> str:
    return f"{SIGLAS[d.month - 1]}{d.year % 100:02d}"


def nome_aba_base(d: date) -> str:
    return f"{SIGLAS[d.month - 1]} {d.year}"


def ruido(a: float) -> float:
    return random.uniform(-a, a)


# ---------------------------------------------------------------------------------------------------------------------
# As sete obras (roteiro de cada uma em nós mensais: atraso em meses, IDC, desvio de caixa e indicadores de canteiro)
# ---------------------------------------------------------------------------------------------------------------------
@dataclass
class Obra:
    obra: str
    bairro: str
    codigo: str
    inicio: date          # mês em que a execução começa
    prazo: int            # meses de obra (início ao término previsto, inclusive)
    folga_entrega: int    # meses entre o término previsto e a entrega contratual
    orc0: float           # orçamento projetado em set/25
    m2: int
    unidades: int
    tipo: list            # [(mês "AAAA-MM" a partir do qual vale, tipo do orçamento)]
    gestor: str
    atraso: dict          # nós do atraso de base (meses), sem a chuva
    recupera: float       # quanto do atraso da chuva a obra recupera por mês, passada a quadra
    idc: dict             # nós do IDC acumulado
    caixa: dict           # nós do desvio de caixa do mês (desembolsado ÷ previsto − 1)
    ppc: dict
    inspecao: dict
    nr18: dict
    entregas: dict
    zona_norte: bool = False
    acidentes: dict = field(default_factory=dict)   # {"AAAA-MM": (acidentes, dias perdidos)}
    saldo: dict = field(default_factory=dict)       # saldo de contrato: {"AAAA-MM": valor} (nós; 0 = sem)
    adiant_mat: dict = field(default_factory=dict)  # {"AAAA-MM": fração do orçamento}
    adiant_serv: dict = field(default_factory=dict)
    receita: tuple = ()                             # (mês inicial "AAAA-MM", valor inicial, acréscimo por mês)
    divida: dict = field(default_factory=dict)      # {"AAAA-MM": (descrição, valor)}
    irregulares: dict = field(default_factory=dict)
    equipamentos: list = field(default_factory=list)
    injetar: dict = field(default_factory=dict)     # falhas de preenchimento (item 9 da régua e fConferencia)

    @property
    def termino(self) -> date:
        return add_mes(self.inicio, self.prazo - 1)

    @property
    def entrega(self) -> date:
        return add_mes(self.termino, self.folga_entrega)

    def tipo_em(self, d: date) -> str:
        return [t for k, t in self.tipo if chave(d) >= k][-1]


EQUIP_PADRAO = [("Betoneira 400 L", 9_800.0), ("Guincho de coluna", 38_500.0), ("Andaime fachadeiro (lote)", 64_000.0),
                ("Serra circular de bancada", 6_400.0), ("Gerador 40 kVA", 52_000.0)]

OBRAS = [
    Obra("CAJUEIRO", "Candelária", "CJ", date(2023, 10, 1), 39, 2, 21_400_000, 6_900, 72,
         [("2000-01", "Orçamento final")], "Eng. P. Medeiros",
         atraso={"2025-09": 0.4, "2026-03": 0.5, "2026-09": 0.6}, recupera=0.4,
         idc={"2025-09": 1.012, "2026-03": 1.017, "2026-09": 1.015},
         caixa={"2025-09": 0.04, "2026-01": -0.06, "2026-05": 0.07, "2026-09": -0.05},
         ppc={"2024-10": 0.82, "2026-09": 0.86}, inspecao={"2024-10": 0.95, "2026-09": 0.97},
         nr18={"2024-10": 0.95, "2026-09": 0.97}, entregas={"2024-10": 0.91, "2026-09": 0.93},
         receita=("2025-09", 382_000.0, 17_500.0),
         equipamentos=EQUIP_PADRAO[:3] + [("Plataforma elevatória", 71_000.0)],
         injetar={}),
    Obra("PONTA NEGRA", "Ponta Negra", "PN", date(2024, 11, 1), 30, 2, 38_600_000, 9_800, 56,
         [("2000-01", "Orçamento executivo")], "Eng.ª L. Bezerra",
         atraso={"2025-09": 0.3, "2026-09": 0.3}, recupera=0.35,
         idc={"2025-09": 1.004, "2026-02": 1.009, "2026-09": 1.011},
         caixa={"2025-09": 0.05, "2025-12": -0.04, "2026-04": 0.08, "2026-09": 0.03},
         ppc={"2024-10": 0.84, "2026-09": 0.87}, inspecao={"2024-10": 0.96, "2026-09": 0.97},
         nr18={"2024-10": 0.96, "2026-09": 0.97}, entregas={"2024-10": 0.92, "2026-09": 0.94},
         saldo={"2025-12": 0, "2026-01": 140_000.0, "2026-09": 60_000.0},
         adiant_mat={"2025-10": 0.006, "2025-11": 0.006, "2025-12": 0.004},
         adiant_serv={"2026-03": 0.0025, "2026-04": 0.0025, "2026-05": 0.0025},
         receita=("2026-03", 64_000.0, 26_000.0),
         equipamentos=EQUIP_PADRAO + [("Grua ascensional (valor residual)", 186_000.0)],
         injetar={}),
    Obra("DUNAS", "Capim Macio", "DN", date(2025, 3, 1), 28, 2, 27_900_000, 8_400, 96,
         [("2000-01", "Orçamento executivo")], "Eng. R. Fontes",
         atraso={"2025-09": 1.3, "2025-10": 1.2, "2026-02": 0.7, "2026-09": 0.4}, recupera=0.4,
         idc={"2025-09": 0.960, "2025-10": 0.962, "2026-01": 0.972, "2026-04": 0.982, "2026-09": 0.988},
         caixa={"2025-09": 0.06, "2025-11": 0.12, "2026-01": 0.06, "2026-04": 0.04, "2026-09": -0.03},
         ppc={"2024-10": 0.66, "2025-10": 0.74, "2026-04": 0.82, "2026-09": 0.85},
         inspecao={"2024-10": 0.86, "2025-10": 0.905, "2026-04": 0.955, "2026-09": 0.96},
         nr18={"2024-10": 0.86, "2025-10": 0.89, "2026-03": 0.96, "2026-09": 0.97},
         entregas={"2024-10": 0.78, "2025-10": 0.81, "2026-04": 0.91, "2026-09": 0.93},
         saldo={"2025-09": 210_000.0, "2026-09": 30_000.0},
         adiant_mat={"2026-01": 0.005, "2026-02": 0.005, "2026-03": 0.003},
         divida={"2026-02": ("Aço CA-50 cedido pela PONTA NEGRA (2,1 t)", 41_700.0),
                 "2026-03": ("Aço CA-50 cedido pela PONTA NEGRA (2,1 t)", 41_700.0)},
         equipamentos=EQUIP_PADRAO[:4],
         injetar={"justificativa_em_branco": ["2025-11"]}),
    Obra("POTENGI", "Potengi", "PT", date(2025, 2, 1), 29, 2, 18_700_000, 7_100, 144,
         [("2000-01", "Orçamento executivo")], "Eng. T. Galvão", zona_norte=True,
         atraso={"2025-09": 0.8, "2026-02": 1.0, "2026-03": 1.15, "2026-06": 1.9, "2026-07": 2.1, "2026-09": 2.2},
         recupera=0.0,
         idc={"2025-09": 0.992, "2026-09": 0.985},
         caixa={"2025-09": -0.04, "2026-02": -0.06, "2026-04": -0.16, "2026-06": -0.22, "2026-07": -0.18,
                "2026-09": -0.09},
         ppc={"2024-10": 0.74, "2026-02": 0.72, "2026-05": 0.63, "2026-07": 0.66, "2026-09": 0.70},
         inspecao={"2024-10": 0.92, "2026-09": 0.91}, nr18={"2024-10": 0.90, "2026-09": 0.90},
         entregas={"2024-10": 0.82, "2026-01": 0.79, "2026-04": 0.68, "2026-07": 0.70, "2026-09": 0.76},
         acidentes={"2026-04": (1, 18)},
         irregulares={"2025-12": ("NF 30914 · lançada na POTENGI, é da REDINHA", 9_870.0)},
         equipamentos=EQUIP_PADRAO[:2] + [("Compactador de solo", 14_200.0)],
         injetar={"justificativa_em_branco": ["2026-06"]}),
    Obra("REDINHA", "Redinha", "RD", date(2024, 8, 1), 32, 2, 24_300_000, 7_600, 88,
         [("2000-01", "Orçamento executivo")], "Eng.ª M. Queiroz", zona_norte=True,
         atraso={"2025-09": 0.5, "2026-09": 0.8}, recupera=0.2,
         idc={"2025-09": 0.974, "2025-12": 0.962, "2026-03": 0.945, "2026-06": 0.917, "2026-08": 0.894,
              "2026-09": 0.889},
         caixa={"2025-09": 0.08, "2025-12": 0.11, "2026-03": 0.15, "2026-05": 0.21, "2026-06": 0.26, "2026-09": 0.19},
         ppc={"2024-10": 0.78, "2026-09": 0.74}, inspecao={"2024-10": 0.93, "2026-03": 0.925, "2026-09": 0.89},
         nr18={"2024-10": 0.93, "2026-09": 0.90}, entregas={"2024-10": 0.88, "2026-09": 0.85},
         acidentes={"2025-11": (1, 6), "2026-07": (1, 12)},
         adiant_mat={"2026-08": 0.004, "2026-09": 0.004},
         receita=("2026-06", 22_000.0, 11_500.0),
         divida={"2026-04": ("Cimento CP-II (180 sacos) cedido pela POTENGI", 18_400.0),
                 "2026-05": ("Cimento CP-II (180 sacos) cedido pela POTENGI", 18_400.0),
                 "2026-06": ("Cimento CP-II (180 sacos) cedido pela POTENGI", 18_400.0)},
         equipamentos=EQUIP_PADRAO[:4],
         injetar={"resultado_digitado": {"2026-05": 85_000.0}}),
    Obra("JANGADA", "Areia Preta", "JG", date(2025, 6, 1), 32, 2, 33_200_000, 9_200, 64,
         [("2000-01", "Orçamento preliminar"), ("2026-04", "Orçamento executivo")], "Eng. A. Lucena",
         atraso={"2025-09": 0.5, "2026-09": 0.6}, recupera=0.3,
         idc={"2025-09": 0.972, "2026-02": 0.966, "2026-06": 0.968, "2026-09": 0.964},
         caixa={"2025-09": 0.06, "2025-12": -0.07, "2026-03": 0.05, "2026-06": -0.07, "2026-09": 0.05},
         ppc={"2024-10": 0.77, "2026-09": 0.78}, inspecao={"2024-10": 0.935, "2026-09": 0.94},
         nr18={"2024-10": 0.90, "2026-09": 0.92}, entregas={"2024-10": 0.87, "2026-09": 0.88},
         acidentes={"2026-02": (1, 3)},
         adiant_serv={"2026-07": 0.002, "2026-08": 0.002, "2026-09": 0.002},
         irregulares={"2026-06": ("NF 88213 · lançada na JANGADA, é da TIROL", 23_150.0)},
         equipamentos=EQUIP_PADRAO[:2] + [("Bomba de concreto (valor residual)", 92_000.0)],
         injetar={"previsto_digitado": ["2026-03"]}),
    Obra("TIROL", "Tirol", "TR", date(2026, 2, 1), 34, 2, 29_500_000, 8_300, 80,
         [("2000-01", "Orçamento legal")], "Eng.ª S. Varela",
         atraso={"2026-02": 0.1, "2026-09": 0.2}, recupera=0.3,
         idc={"2026-02": 1.010, "2026-05": 1.004, "2026-09": 1.008},
         caixa={"2026-02": 0.06, "2026-04": -0.05, "2026-09": 0.04},
         ppc={"2026-02": 0.78, "2026-09": 0.83}, inspecao={"2026-02": 0.93, "2026-09": 0.95},
         nr18={"2026-02": 0.88, "2026-05": 0.95, "2026-09": 0.96}, entregas={"2026-02": 0.86, "2026-09": 0.92},
         equipamentos=EQUIP_PADRAO[:2] + [("Container escritório (2 un.)", 28_000.0)],
         injetar={"mes_desatualizado": ["2026-04"]}),
]

CLIENTES = ["A. Cavalcanti", "B. Freire", "C. Bezerra", "D. Pinheiro", "E. Maia", "F. Rocha", "G. Barreto",
            "H. Teixeira", "I. Moura", "J. Viana", "K. Saraiva", "L. Fernandes", "M. Albuquerque", "N. Siqueira"]
PERSONALIZACOES = ["troca de piso da sala", "bancada da cozinha em quartzo", "planta com home office",
                   "ponto extra de ar-condicionado", "revestimento da varanda", "fechamento da cozinha",
                   "box e metais de padrão superior", "iluminação embutida"]
ESTOQUE_ITENS = {
    "mobilizacao": ["Material de canteiro (tapumes e contêineres)", "Cimento e agregados", "Aço CA-50 para fundação"],
    "fundacao": ["Aço CA-50 para blocos e baldrame", "Cimento e agregados", "Madeira para formas"],
    "estrutura": ["Aço CA-50/60 para lajes", "Formas e escoramento", "Blocos cerâmicos", "Material hidrossanitário"],
    "alvenaria": ["Blocos cerâmicos", "Argamassa industrializada", "Material hidrossanitário",
                  "Material elétrico (fios e cabos)", "Revestimento cerâmico (compra antecipada)"],
    "acabamento": ["Revestimento cerâmico interno", "Porcelanato das áreas comuns", "Louças e metais",
                   "Material elétrico (fios e cabos)", "Tinta acrílica", "Material para impermeabilização"],
    "entrega": ["Louças e metais de reposição", "Tinta para retoques"],
}
CONTAS_ITENS = {
    "mobilizacao": ["Medição de terraplenagem", "Locação de equipamentos", "Encargos e INSS do mês"],
    "fundacao": ["Medição de fundação (estacas hélice)", "Concreto usinado dos blocos", "Encargos e INSS do mês"],
    "estrutura": ["Medição da empreiteira de estrutura", "Concreto usinado da laje", "Aluguel de escoramento",
                  "Encargos e INSS do mês"],
    "alvenaria": ["Medição da empreiteira de alvenaria", "Medição de instalações hidráulicas",
                  "Medição de instalações elétricas", "Encargos e INSS do mês"],
    "acabamento": ["Medição de reboco e gesso", "Medição de pintura", "Medição de revestimento",
                   "Medição de instalações elétricas", "Encargos e INSS do mês"],
    "entrega": ["Medição de pintura final", "Limpeza final e entrega", "Encargos e INSS do mês"],
}
JUSTIF_ACIMA = ["Compra antecipada de material para travar preço com o fornecedor.",
                "Medição acumulada de duas parcelas da empreiteira no mesmo mês.",
                "Aditivo de serviço não previsto no orçamento.",
                "Aumento de efetivo para recuperar o cronograma.",
                "Reajuste de contrato de mão de obra pelo INCC."]
JUSTIF_ABAIXO = ["Medição da empreiteira postergada para o mês seguinte.",
                 "Chuvas reduziram a produção e a medição do mês.",
                 "Entrega de material remarcada pelo fornecedor.",
                 "Concretagem adiada por chuva; pagamento passa para o mês seguinte."]


# ---------------------------------------------------------------------------------------------------------------------
# Série de cada obra
# ---------------------------------------------------------------------------------------------------------------------
@dataclass
class Mes:
    mes: date
    t: int                 # meses desde o início (1 = mês de início)
    prev_acum: float       # previsto acumulado (linha de base)
    exec_acum: float
    atraso: float
    fase: str
    chuva: int
    valores: dict = field(default_factory=dict)


def dias_chuva(o: Obra, d: date) -> int:
    extra = 1 if (o.zona_norte and d.month in (3, 4, 5, 6, 7)) else 0
    return max(0, CHUVA[d.month] + extra + random.choice((-1, 0, 0, 1)) * (CHUVA[d.month] > 1))


def serie_fisica(o: Obra) -> list[Mes]:
    """Previsto e executado acumulados de cada mês, do início até o corte. O executado sai do atraso: a obra está, em
    cada mês, no ponto da curva prevista de `atraso` meses atrás (cronograma agregado). A chuva soma atraso na quadra
    chuvosa, conforme a fase, e a obra recupera parte dele depois."""
    saida, e_ant, chuva_atraso = [], 0.0, 0.0
    for d in meses(o.inicio, CORTE):
        t = meses_entre(o.inicio, d) + 1
        p = curva_s(t / o.prazo)
        fs = fase(e_ant if e_ant else p)
        dc = dias_chuva(o, d)
        if d >= date(2025, 9, 1):
            chuva_atraso += dc * SENS_CHUVA[fs] / 21
            if dc <= 2:
                chuva_atraso *= (1 - o.recupera)
        a = interp(o.atraso, d) + chuva_atraso
        e = max(curva_s((t - a) / o.prazo), e_ant + 0.0012)
        saida.append(Mes(d, t, p, e, a, fase(e), dc))
        e_ant = e
    return saida


def r2(v: float) -> float:
    return round(v + 0.0, 2)


def valores_mensais(o: Obra, serie: list[Mes]) -> list[Mes]:
    """Números da Análise Crítica de cada mês (de ago/25, para o desembolso do mês de set/25, até o corte)."""
    ant = None
    saida = []
    for s in serie:
        if s.mes < add_mes(INICIO_AC, -1):
            continue
        k = chave(s.mes)
        v = s.valores
        orc = r2(o.orc0 * (1 + INCC_MES * meses_entre(INICIO_AC, s.mes)))
        f35 = round(s.prev_acum, 4)
        f37 = round(s.exec_acum, 4)
        if k in o.injetar.get("previsto_digitado", []):
            f35 = round(ant.valores["F35"] - 0.004, 4)      # reprogramação digitada sem ajustar o histórico
        f35_ant = ant.valores["F35_base"] if ant else round(curva_s((s.t - 1) / o.prazo), 4)
        f37_ant = ant.valores["F37"] if ant else round(curva_s((s.t - 1 - s.atraso) / o.prazo), 4) if s.t > 1 else 0.0
        v.update(orc=orc, tipo=o.tipo_em(s.mes), F35=f35, F35_base=round(s.prev_acum, 4), F37=f37,
                 F31=round(round(s.prev_acum, 4) - f35_ant, 4), F33=round(f37 - f37_ant, 4))
        va = r2(f37 * orc)
        idc = interp(o.idc, s.mes)
        ca = va / idc
        est_frac = ESTOQUE_FRAC[s.fase] * random.uniform(0.8, 1.2)
        estoque = r2(orc * est_frac) if s.t > 1 else r2(orc * 0.001)
        adm = r2(orc * o.adiant_mat.get(k, 0.0))
        ads = r2(orc * o.adiant_serv.get(k, 0.0))
        saldo = r2(max(0.0, interp(o.saldo, s.mes))) if o.saldo else 0.0
        equip = r2(sum(x for _n, x in o.equipamentos))
        receita = 0.0
        if o.receita and k >= o.receita[0]:
            receita = r2(o.receita[1] + o.receita[2] * meses_entre(date(int(o.receita[0][:4]), int(o.receita[0][5:]), 1),
                                                                   s.mes))
        exec_rs = max(v["F33"], 0.002) * orc
        contas = r2(max(20_000.0, exec_rs * random.uniform(0.16, 0.26)))
        div = o.divida.get(k)
        divida = r2(div[1]) if div else 0.0
        irr = o.irregulares.get(k)
        irregular = r2(irr[1]) if irr else 0.0
        corrigido = r2(ca - contas - divida + estoque + adm + ads + receita)
        if ant is not None:
            minimo = ant.valores["corrigido"] + orc * 0.004
            if corrigido < minimo:                         # o desembolso acumulado não anda para trás
                estoque = r2(estoque + (minimo - corrigido))
                corrigido = r2(ca - contas - divida + estoque + adm + ads + receita)
        f55 = r2(corrigido + irregular)
        v.update(va=va, ca=r2(ca), idc_alvo=idc, estoque=estoque, adiant_mat=adm, adiant_serv=ads, saldo=saldo,
                 equipamentos=equip, receita=receita, contas=contas, divida=divida, irregulares=irregular,
                 corrigido=corrigido, F55=f55)
        if ant is not None or s.t == 1:
            desemb_mes = r2(f55 - (ant.valores["F55"] if ant is not None else 0.0))
            desvio = interp(o.caixa, s.mes) + ruido(0.02)
            previsto = r2(desemb_mes / (1 + desvio))
            exigida = abs(desemb_mes / previsto - 1) > 0.10
            justificativa = None
            if exigida and k not in o.injetar.get("justificativa_em_branco", []):
                justificativa = random.choice(JUSTIF_ACIMA if desemb_mes > previsto else JUSTIF_ABAIXO)
            v.update(J136=previsto, J137=desemb_mes, justificativa=justificativa)
        resultado_regra = r2(va + saldo + ads + adm + estoque + equip + receita - contas - corrigido - divida)
        digitado = o.injetar.get("resultado_digitado", {}).get(k)
        v.update(resultado_regra=resultado_regra,
                 resultado=r2(resultado_regra + digitado) if digitado is not None else resultado_regra,
                 resultado_digitado=digitado is not None)
        saida.append(s)
        ant = s
    return saida


def indicadores_base(o: Obra, serie: list[Mes]) -> list[dict]:
    """Uma linha por mês da base complementar 2.0 (de out/24, ou do início da obra, até o corte)."""
    linhas = []
    for s in serie:
        if s.mes < INICIO_BASE:
            continue
        k = chave(s.mes)
        efetivo = max(14, round(o.orc0 / 1e6 * EFETIVO_POR_MI[s.fase] * random.uniform(0.94, 1.06)))
        hht = round(efetivo * (22 - 0.5 * s.chuva) * 8.8 * random.uniform(0.97, 1.03))
        verificados = max(8, round(efetivo * random.uniform(0.45, 0.65)))
        aprovados = min(verificados, round(verificados * min(1.0, interp(o.inspecao, s.mes) + ruido(0.015))))
        aplicaveis = random.randint(118, 152)
        conformes = min(aplicaveis, round(aplicaveis * min(1.0, interp(o.nr18, s.mes) + ruido(0.012))))
        previstos = max(10, round(efetivo * random.uniform(0.32, 0.48)))
        no_prazo = min(previstos, round(previstos * min(1.0, interp(o.entregas, s.mes) + ruido(0.02))))
        acid, dias = o.acidentes.get(k, (0, 0))
        linhas.append({"obra": o.obra, "mes": s.mes.isoformat(), "aba": nome_aba_base(s.mes), "gestor": o.gestor,
                       "PPC": round(min(0.98, interp(o.ppc, s.mes) + ruido(0.025)), 2),
                       "ServicosVerificados": verificados, "AprovadosPrimeira": aprovados, "HHT": hht,
                       "Acidentes": acid, "DiasPerdidos": dias, "NR18Aplicaveis": aplicaveis,
                       "NR18Conformes": conformes, "PedidosPrevistos": previstos, "PedidosNoPrazo": no_prazo,
                       "DiasChuva": s.chuva, "efetivo": efetivo})
    return linhas


# ---------------------------------------------------------------------------------------------------------------------
# Planilha Análise Crítica (layout lido por rótulo em fxLerAba.m e pipeline/extrair_analise_critica.py)
# ---------------------------------------------------------------------------------------------------------------------
VERDE = PatternFill("solid", fgColor="C6EFCE")       # célula editável
LARANJA = PatternFill("solid", fgColor="FFE699")     # resultado (não editar)
TABELA = PatternFill("solid", fgColor="DDEBF7")      # tabela anexada
PCT_FILL = PatternFill("solid", fgColor="FCE4D6")    # % relacionada
NEGRITO = Font(bold=True)
TITULO = Font(bold=True, size=14)
FINO = Side(style="thin", color="BFBFBF")
CAIXA = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
RS = '#,##0.00;[Red]-#,##0.00'
# Data de mês: formato EMBUTIDO do Excel (numFmtId 17, mmm-yy). O Excel.Workbook do Power Query só entrega como data a
# célula com formato de data embutido; com formato personalizado ("mmm/yy") entrega o número de série. O openpyxl grava
# "mmm-yy" como o 17; o LibreOffice, ao recalcular, regrava tudo como personalizado (datas_embutidas desfaz).
DATA_MES = "mmm-yy"
DATA_EMBUTIDA = 17
PCT = '0.00%'

# (linha do título, coluna da descrição, rótulo, linha do total, quantas linhas)
TABELAS_AC = [
    (63, 2, "Saldo de contrato", 74, 10), (63, 7, "Adiantamento de materiais não estocados", 74, 10),
    (77, 2, "Adiantamento de serviços não medidos", 88, 10), (77, 7, "Estoque de materiais", 88, 10),
    (91, 2, "Equipamentos comprados", 102, 10), (91, 7, "Receitas de personalizações", 102, 10),
    (105, 2, "Contas a pagar dos medidos", 115, 9), (105, 7, "Itens apontados irregularmente no relatório", 115, 9),
    (118, 2, "Dívida da obra a ser paga", 129, 10),
]
SIGNIFICADOS = [
    ("Crédito do orçamento", "Quanto a obra podia ter gasto pelo que executou: orçamento × executado acumulado."),
    ("Saldo de contrato", "Desconto obtido num contrato e ainda não realizado; abatido à medida que o serviço é medido."),
    ("Adiantamento de serviços não medidos", "Sinal pago de contrato ainda não medido (ex.: elevadores)."),
    ("Adiantamento de materiais não estocados", "Material pago que ainda não chegou à obra (ex.: esquadrias)."),
    ("Estoque de materiais", "Material no almoxarifado, ainda não aplicado, informado pelo almoxarife."),
    ("Equipamentos (valor de mercado)", "Equipamento comprado pela obra e reaproveitável no fim, pelo valor de mercado."),
    ("Receita de personalizações", "Personalizações pagas pelos clientes, acumuladas até o fim da obra."),
    ("Contas a pagar dos medidos", "Serviço já medido e ainda a pagar (relatório de contas a pagar)."),
    ("Valor desembolsado pela obra", "Desembolso acumulado desde o início, conferido e sem os itens apontados na obra "
                                     "errada."),
    ("Dívida da obra a ser paga", "Material de outra obra da carteira usado nesta: custo da obra, desconta do "
                                  "resultado (versão 01 da planilha)."),
]


def _celula(ws, ref: str, valor, fill=None, fmt=None, bold=False):
    c = ws[ref]
    c.value = valor
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if bold:
        c.font = NEGRITO
    return c


def escrever_aba_ac(ws, o: Obra, s: Mes, aba_ant: str | None):
    v = s.valores
    k = chave(s.mes)
    for col, larg in (("A", 3), ("B", 46), ("C", 30), ("D", 18), ("E", 18), ("F", 18), ("G", 44), ("H", 12),
                      ("I", 12), ("J", 18)):
        ws.column_dimensions[col].width = larg
    _celula(ws, "B2", "ANÁLISE DE DESEMPENHO ECONÔMICO DA OBRA").font = TITULO
    _celula(ws, "B3", "Construtora fictícia de Natal/RN · portfólio Vizu Dados · dados inventados")
    _celula(ws, "B4", "Versão 01 · revisão de 01/09/2025: a dívida da obra passa a descontar do resultado")
    _celula(ws, "B18", "Célula editável", VERDE)
    _celula(ws, "B19", "Resultado (não editar)", LARANJA)
    _celula(ws, "B20", "Tabela anexada", TABELA)
    _celula(ws, "B21", "% relacionada", PCT_FILL)
    mes_celula = add_mes(s.mes, -1) if k in o.injetar.get("mes_desatualizado", []) else s.mes
    passos = [
        (24, "1º passo: Projeto (definir o empreendimento)", "D24", o.obra, VERDE, None),
        (26, "2º passo: Mês de análise (definir o mês de referência)", "E26", datetime(mes_celula.year, mes_celula.month, 1),
         VERDE, DATA_MES),
        (28, "3º passo: Orçamento da obra projetado (com INCC no a realizar)", "J28", v["orc"], VERDE, RS),
        (29, "Tipo do orçamento", "J29", v["tipo"], VERDE, None),
        (31, "4º passo: Previsto do mês (incluir o % do cronograma)", "F31", v["F31"], VERDE, PCT),
        (33, "5º passo: Executado do mês (incluir o % medido)", "F33", v["F33"], VERDE, PCT),
        (35, "6º passo: Previsto acumulado do mês", "F35", v["F35"], VERDE, PCT),
        (37, "7º passo: Executado acumulado do mês", "F37", v["F37"], VERDE, PCT),
        (39, "8º passo: Crédito do orçamento (executado acumulado × orçamento)", "F39", "=F37*J28", LARANJA, RS),
        (41, "8º passo (cont.): Saldo de contrato (tabela)", "F41", "=E74", LARANJA, RS),
        (43, "9º passo: Adiantamento de serviços não medidos (tabela)", "F43", "=E88", LARANJA, RS),
        (45, "10º passo: Adiantamento de materiais não estocados (tabela)", "F45", "=J74", LARANJA, RS),
        (47, "11º passo: Estoque de materiais (almoxarifado)", "F47", "=J88", LARANJA, RS),
        (49, "12º passo: Equipamentos comprados (valor de mercado)", "F49", "=E102", LARANJA, RS),
        (51, "13º passo: Receitas de personalizações (tabela)", "F51", "=J102", LARANJA, RS),
        (53, "14º passo: Contas a pagar dos medidos (relatório a pagar)", "F53", "=E115", LARANJA, RS),
        (55, "15º passo: Valor desembolsado pela obra (relatório conferido)", "F55", v["F55"], VERDE, RS),
        (56, "Valor corrigido (sem os itens apontados irregularmente)", "F56", "=F55-J115", LARANJA, RS),
        (58, "16º passo: Dívida da obra a ser paga (outra obra da carteira)", "F58", "=E129", LARANJA, RS),
    ]
    for linha, rotulo, ref, valor, fill, fmt in passos:
        _celula(ws, f"B{linha}", rotulo)
        _celula(ws, ref, valor, fill, fmt)

    # tabelas anexadas
    linhas = tabelas_do_mes(o, s)
    for tl, col, rotulo, ltotal, n in TABELAS_AC:
        cd = openpyxl.utils.get_column_letter(col)
        cv = openpyxl.utils.get_column_letter(col + 3)
        _celula(ws, f"{cd}{tl}", rotulo, TABELA, bold=True)
        for i, (desc, val) in enumerate(linhas.get(rotulo, [])[:n]):
            _celula(ws, f"{cd}{tl + 1 + i}", desc)
            _celula(ws, f"{cv}{tl + 1 + i}", val, fmt=RS)
        _celula(ws, f"{cd}{ltotal}", "Total", bold=True)
        _celula(ws, f"{cv}{ltotal}", f"=SUM({cv}{tl + 1}:{cv}{ltotal - 1})", LARANJA, RS, True)

    # bloco de resultado (área de impressão)
    _celula(ws, "B132", "RESULTADO PARA UTILIZAR NA APRESENTAÇÃO").font = TITULO
    info = [(136, "Projeto", "=D24", None), (137, "Orçamento", "=J28", RS), (138, "Tipo do orçamento", "=J29", None),
            (139, "Mês", "=E26", DATA_MES), (140, "% previsto do mês", "=F31", PCT), (141, "% executado do mês", "=F33", PCT),
            (142, "% previsto acumulado", "=F35", PCT), (143, "% executado acumulado", "=F37", PCT)]
    for linha, rotulo, f, fmt in info:
        _celula(ws, f"C{linha}", rotulo)
        _celula(ws, f"D{linha}", f, LARANJA, fmt)
    itens = [(145, "Crédito do orçamento", "=F39"), (146, "Saldo de contrato", "=F41"),
             (147, "Adiantamento de serviços não medidos", "=F43"), (148, "Adiantamento de materiais não estocados", "=F45"),
             (149, "Estoque de materiais", "=F47"), (150, "Equipamentos (valor de mercado)", "=F49"),
             (151, "Receita de personalizações", "=F51"), (152, "Contas a pagar dos medidos", "=F53"),
             (153, "Valor desembolsado (corrigido)", "=F56"), (154, "Dívida da obra a ser paga", "=F58")]
    for i, (linha, rotulo, f) in enumerate(itens, 1):
        _celula(ws, f"B{linha}", f"{i:02d}")
        _celula(ws, f"C{linha}", rotulo)
        _celula(ws, f"D{linha}", f, LARANJA, RS)
        _celula(ws, f"E{linha}", f"=D{linha}/$D$137", PCT_FILL, PCT)
    _celula(ws, "C155", "Resultado mensal da obra", bold=True)
    if v["resultado_digitado"]:
        _celula(ws, "D155", v["resultado"], LARANJA, RS, True)      # digitado por cima da fórmula
    else:
        _celula(ws, "D155", "=D145+D146+D147+D148+D149+D150+D151-D152-D153-D154", LARANJA, RS, True)
    _celula(ws, "E155", "=D155/D137", LARANJA, PCT, True)
    _celula(ws, "C156", "Resultado do mês anterior")
    if aba_ant:
        _celula(ws, "D156", f"='{aba_ant}'!D155", LARANJA, RS)
        _celula(ws, "E156", f"='{aba_ant}'!E155", LARANJA, PCT)
    # tabela informativa do desembolso do mês
    _celula(ws, "G135", "Tabela informativa do desembolso do mês", bold=True)
    _celula(ws, "G136", "Valor previsto para o mês")
    _celula(ws, "J136", v.get("J136"), VERDE, RS)
    _celula(ws, "G137", "Valor desembolsado no mês")
    _celula(ws, "J137", v.get("J137"), VERDE, RS)
    _celula(ws, "G138", "Diferença entre o previsto e o desembolsado")
    _celula(ws, "J138", "=J137-J136", LARANJA, RS)
    _celula(ws, "G139", "Justificativa para diferença acima de 10% (obrigatória)")
    if v.get("justificativa"):
        _celula(ws, "G140", v["justificativa"], VERDE)
        ws["G140"].alignment = Alignment(wrap_text=True)
    # significado de cada item (texto de ajuda que o painel lê)
    _celula(ws, "B160", "Significado de cada item", bold=True)
    for i, (rotulo, texto) in enumerate(SIGNIFICADOS):
        _celula(ws, f"B{162 + 2 * i}", f"{rotulo}:", bold=True)
        _celula(ws, f"B{163 + 2 * i}", texto)
    ws.print_area = "A132:J157"


def tabelas_do_mes(o: Obra, s: Mes) -> dict[str, list[tuple[str, float]]]:
    """Linhas das tabelas anexadas do mês (descrição, valor), somando o que o roteiro pôs em cada item."""
    v = s.valores
    rnd = random.Random(f"{o.obra}{s.mes}")
    t: dict[str, list] = {}
    if v["saldo"]:
        t["Saldo de contrato"] = [("Desconto no contrato de esquadrias (abatido na medição)", v["saldo"])]
    if v["adiant_mat"]:
        t["Adiantamento de materiais não estocados"] = [("Esquadrias de alumínio · sinal de fabricação", v["adiant_mat"])]
    if v["adiant_serv"]:
        t["Adiantamento de serviços não medidos"] = [("Elevadores · sinal de contrato", v["adiant_serv"])]
    t["Estoque de materiais"] = dividir(v["estoque"], ESTOQUE_ITENS[s.fase], rnd)
    t["Equipamentos comprados"] = [(n, x) for n, x in o.equipamentos]
    if v["receita"]:
        n = min(10, 2 + meses_entre(INICIO_AC, s.mes) // 2)
        fixo = random.Random(o.obra)                       # as mesmas unidades e clientes em todos os meses
        rotulos = []
        for i in range(10):
            unid = 100 * fixo.randint(2, 18) + fixo.randint(1, 4)
            rotulos.append(f"Unid. {unid} · {PERSONALIZACOES[i % len(PERSONALIZACOES)]} "
                           f"(cliente {CLIENTES[(i + len(o.obra)) % len(CLIENTES)]})")
        t["Receitas de personalizações"] = dividir(v["receita"], rotulos[:n], random.Random(o.obra + "R"))
    nfs = [f"{d} · NF {rnd.randint(1000, 99999)}" for d in CONTAS_ITENS[s.fase]]
    t["Contas a pagar dos medidos"] = dividir(v["contas"], nfs, rnd)
    k = chave(s.mes)
    if k in o.irregulares:
        t["Itens apontados irregularmente no relatório"] = [(o.irregulares[k][0], v["irregulares"])]
    if k in o.divida:
        t["Dívida da obra a ser paga"] = [(o.divida[k][0], v["divida"])]
    return t


def dividir(total: float, nomes: list[str], rnd: random.Random) -> list[tuple[str, float]]:
    if total <= 0 or not nomes:
        return []
    pesos = [rnd.uniform(0.6, 1.4) for _ in nomes]
    s = sum(pesos)
    partes = [round(total * p / s, 2) for p in pesos]
    partes[-1] = round(total - sum(partes[:-1]), 2)
    return list(zip(nomes, partes))


def leia_me(wb, linhas: list[str]):
    ws = wb.create_sheet("LEIA-ME", 0)
    ws.column_dimensions["B"].width = 110
    for i, t in enumerate(linhas, 2):
        ws.cell(i, 2, t).font = TITULO if i == 2 else Font()


def gravar_ac(o: Obra, serie: list[Mes], destino: Path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    leia_me(wb, [f"Análise Crítica · {o.obra}", "Dados fictícios: construtora inventada de Natal/RN, portfólio Vizu Dados.",
                 "Uma aba por mês (MMMAA). Os campos são lidos pelo rótulo da coluna B.",
                 "Versão 01 da planilha: a dívida da obra (item 10) desconta do resultado."])
    aba_ant = None
    for s in serie:
        if s.mes < INICIO_AC:
            continue
        nome = nome_aba_ac(s.mes)
        ws = wb.create_sheet(nome)
        escrever_aba_ac(ws, o, s, aba_ant)
        aba_ant = nome
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


# ---------------------------------------------------------------------------------------------------------------------
# Base complementar 2.0 (lida por rótulo em stgAbasADO.m: âncora "Versão do modelo" na coluna B, valor na coluna D)
# ---------------------------------------------------------------------------------------------------------------------
SECOES_BASE = [("Planejamento", ["PPC"]), ("Qualidade", ["ServicosVerificados", "AprovadosPrimeira"]),
               ("Segurança", ["HHT", "Acidentes", "DiasPerdidos", "NR18Aplicaveis", "NR18Conformes"]),
               ("Suprimentos", ["PedidosPrevistos", "PedidosNoPrazo"]), ("Clima", ["DiasChuva"])]


def gravar_base(o: Obra, linhas: list[dict], destino: Path):
    rotulo = {c: r for c, r, _f in NOTA["base"]}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    leia_me(wb, [f"Base complementar · {o.obra}", "Dados fictícios: construtora inventada de Natal/RN, portfólio Vizu Dados.",
                 f"Modelo {NOTA['versao_base']}: uma aba por mês (MMM AAAA), um valor por campo na coluna D.",
                 "PPC em fração (0,82 = 82%). HHT, acidentes e dias perdidos pela NBR 14280; itens NR-18 pelo checklist "
                 "do canteiro; pedidos pela programação de suprimentos."])
    for ln in linhas:
        ws = wb.create_sheet(ln["aba"])
        ws.column_dimensions["B"].width = 44
        ws.column_dimensions["D"].width = 16
        ws["B2"] = "BASE COMPLEMENTAR DA OBRA"
        ws["B2"].font = TITULO
        ws["B4"], ws["D4"] = "Versão do modelo", NOTA["versao_base"]
        ws["B5"], ws["D5"] = "Obra", o.obra
        ws["B6"], ws["D6"] = "Gestor da obra", ln["gestor"]
        r = 8
        for secao, campos in SECOES_BASE:
            ws.cell(r, 2, secao).font = NEGRITO
            r += 1
            for c in campos:
                ws.cell(r, 2, rotulo[c])
                cel = ws.cell(r, 4, ln[c])
                cel.fill = VERDE
                cel.number_format = "0%" if c == "PPC" else "#,##0"
                r += 1
            r += 1
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


# ---------------------------------------------------------------------------------------------------------------------
# Cronograma físico-financeiro (lida em stgCrono.m): PREVXREA do corte e uma aba FLUXO-FINACEIRO por obra
# ---------------------------------------------------------------------------------------------------------------------
def gravar_crono(obras: list[tuple[Obra, list[Mes]]], destino: Path) -> list[dict]:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    leia_me(wb, ["Planilha de fluxo financeiro · cronograma físico-financeiro",
                 "Dados fictícios: construtora inventada de Natal/RN, portfólio Vizu Dados.",
                 "PREVXREA: previsto × realizado no mês de corte e o previsto dos meses seguintes, por obra.",
                 "FLUXO-FINACEIRO <estágio> <obra>: a curva prevista (linha de base) mês a mês, do início ao término."])
    ws = wb.create_sheet(f"PREVXREA {SIGLAS[CORTE.month - 1]} {CORTE.year % 100:02d}")
    fim = max(o.termino for o, _s in obras)
    futuros = meses(add_mes(CORTE, 1), fim)
    ws["A1"] = f"PREVISTO X REALIZADO · {SIGLAS[CORTE.month - 1]}/{CORTE.year}"
    ws["A1"].font = TITULO
    cab = ["OBRA", "TIPO DE ORÇAMENTO", "M2 EQUIVALENTE (NBR)", "ORÇ OBRA RESULTADO", "REALIZADO ACUMULADO", "PREVISTO",
           "DESVIO"] + [""] * len(futuros)
    for c, t in enumerate(cab, 1):
        ws.cell(5, c, t).font = NEGRITO
    ws.cell(6, 6, datetime(CORTE.year, CORTE.month, 1)).number_format = DATA_MES
    for i, d in enumerate(futuros):
        ws.cell(6, 8 + i, datetime(d.year, d.month, 1)).number_format = DATA_MES
    curvas = []
    for r, (o, serie) in enumerate(obras, 7):
        ult = serie[-1]
        orc = ult.valores["orc"]
        def custo_mes(d: date) -> float:
            t = meses_entre(o.inicio, d) + 1
            return r2((curva_s(t / o.prazo) - curva_s((t - 1) / o.prazo)) * orc) if 1 <= t <= o.prazo else 0.0
        tipo = ult.valores["tipo"].replace("Orçamento ", "").upper()
        realizado = ult.valores["ca"]                      # custo aplicado até o corte
        ws.cell(r, 1, o.obra)
        ws.cell(r, 2, tipo)
        ws.cell(r, 3, o.m2).number_format = "#,##0"
        ws.cell(r, 4, orc).number_format = RS
        ws.cell(r, 5, realizado).number_format = RS
        ws.cell(r, 6, custo_mes(CORTE)).number_format = RS
        ws.cell(r, 7, r2(ult.valores["J137"] - custo_mes(CORTE))).number_format = RS
        for i, d in enumerate(futuros):
            x = custo_mes(d)
            if x:
                ws.cell(r, 8 + i, x).number_format = RS
        # aba da curva prevista (linha de base): mês na coluna B, custo previsto na coluna E (stgCrono.m)
        estagio = "EX" if "executivo" in ult.valores["tipo"].lower() or "final" in ult.valores["tipo"].lower() else (
            "LE" if "legal" in ult.valores["tipo"].lower() else "PRE")
        wf = wb.create_sheet(f"FLUXO-FINACEIRO {estagio} {o.obra}")
        for c, t in ((2, "MÊS"), (3, "% PREVISTO NO MÊS"), (4, "% PREVISTO ACUMULADO"), (5, "CUSTO PREVISTO")):
            wf.cell(3, c, t).font = NEGRITO
        wf.column_dimensions["B"].width = 12
        wf.column_dimensions["E"].width = 18
        for i, d in enumerate(meses(o.inicio, o.termino)):
            t = i + 1
            p, p0 = curva_s(t / o.prazo), curva_s((t - 1) / o.prazo)
            wf.cell(4 + i, 2, datetime(d.year, d.month, 1)).number_format = DATA_MES
            wf.cell(4 + i, 3, round(p - p0, 6)).number_format = PCT
            wf.cell(4 + i, 4, round(p, 6)).number_format = PCT
            wf.cell(4 + i, 5, custo_mes(d)).number_format = RS
            curvas.append({"obra": o.obra, "mes": d.isoformat(), "previsto_acum": round(p, 6), "custo_previsto": custo_mes(d)})
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return curvas


# ---------------------------------------------------------------------------------------------------------------------
# Cadastro, gabarito e recálculo
# ---------------------------------------------------------------------------------------------------------------------
def gravar_csv(path: Path, linhas: list[dict]):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=list(linhas[0].keys()), delimiter=";")
        w.writeheader()
        w.writerows(linhas)


def datas_embutidas(p: Path) -> int:
    """Depois do LibreOffice: no styles.xml, troca os formatos de data personalizados (numFmtId ≥ 164, que o LibreOffice
    grava até para o formato embutido) pelo embutido 17 (mmm-yy) e apaga a definição deles. Devolve quantos trocou."""
    with zipfile.ZipFile(p) as z:
        itens = [(i, z.read(i.filename)) for i in z.infolist()]
    trocados = 0
    saida = []
    for info, dados in itens:
        if info.filename == "xl/styles.xml":
            x = dados.decode("utf-8")
            datas = {}
            for m in re.finditer(r'<numFmt numFmtId="(\d+)" formatCode="([^"]*)"/>', x):
                cod = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.', "", html.unescape(m.group(2)).lower())
                if re.search(r"[yd]", cod) and not re.search(r"[0#?]", cod):
                    datas[m.group(1)] = m.group(0)
            for nid, defin in datas.items():
                x = x.replace(defin, "")
                x = re.sub(rf'(<xf [^>]*numFmtId=")({nid})(")', rf"\g<1>{DATA_EMBUTIDA}\g<3>", x)
            x = re.sub(r'<numFmts count="\d+"',
                       lambda _m: f'<numFmts count="{len(re.findall("<numFmt ", x))}"', x)
            trocados = len(datas)
            dados = x.encode("utf-8")
        saida.append((info, dados))
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        for info, dados in saida:
            z.writestr(info, dados)
    return trocados


def recalcular(arquivos: list[Path]):
    """Abre e salva cada .xlsx no LibreOffice (headless): grava o valor de cada fórmula, que o Power Query lê. Depois
    devolve às datas o formato embutido (datas_embutidas)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("LibreOffice (soffice) não encontrado: rode com --sem-recalculo ou instale o LibreOffice")
    with tempfile.TemporaryDirectory() as tmp:
        for p in arquivos:
            r = subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(p)],
                               capture_output=True, text=True, timeout=180)
            saida = Path(tmp) / p.name
            if r.returncode or not saida.is_file():
                sys.exit(f"LibreOffice não recalculou {p.name}: {r.stderr[-400:]}")
            shutil.move(str(saida), p)
            datas_embutidas(p)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--sem-recalculo", action="store_true", help="não abre no LibreOffice (fórmulas sem valor salvo)")
    args = ap.parse_args(argv)

    for sub in ("analise-critica", "bases-complementares", "cronograma", "cadastro", "_gabarito"):
        alvo = DADOS / sub
        if alvo.is_dir():
            shutil.rmtree(alvo)
    arquivos, gab_ac, gab_base, por_obra, cadastro = [], [], [], [], []
    for o in OBRAS:
        fisica = serie_fisica(o)
        serie = valores_mensais(o, fisica)
        base = indicadores_base(o, fisica)
        por_obra.append((o, serie))
        ac = DADOS / "analise-critica" / o.obra / f"Analise Critica {o.obra}.xlsx"
        gravar_ac(o, serie, ac)
        bc = DADOS / "bases-complementares" / f"Base complementar - {o.obra}.xlsx"
        gravar_base(o, base, bc)
        arquivos.append(ac)
        for s in serie:
            if s.mes < INICIO_AC:
                continue
            v = s.valores
            gab_ac.append({
                "obra": o.obra, "mes": s.mes.isoformat(), "aba": nome_aba_ac(s.mes), "tipo_orcamento": v["tipo"],
                "orcamento": v["orc"], "prev_mes": v["F31"], "exec_mes": v["F33"], "prev_acum": v["F35"],
                "exec_acum": v["F37"], "credito": v["va"], "saldo_contrato": v["saldo"], "adiant_servicos": v["adiant_serv"],
                "adiant_materiais": v["adiant_mat"], "estoque": v["estoque"], "equipamentos": v["equipamentos"],
                "receita_personalizacao": v["receita"], "contas_pagar": v["contas"], "desembolsado_bruto": v["F55"],
                "itens_irregulares": v["irregulares"], "desembolsado_corrigido": v["corrigido"], "divida": v["divida"],
                "resultado": v["resultado"], "resultado_recalculado": v["resultado_regra"],
                "previsto_mes_rs": v.get("J136"), "desembolsado_mes_rs": v.get("J137"),
                "justificativa": v.get("justificativa"), "atraso_roteiro": round(s.atraso, 4),
                "idc_roteiro": round(v["idc_alvo"], 4), "fase": s.fase, "dias_chuva": s.chuva})
        gab_base += base
        cadastro.append({"obra": o.obra, "codigo": o.codigo, "spe": f"{o.obra} EMPREENDIMENTO SPE LTDA",
                         "nome_crono": o.obra, "fonte_fisico": "analise_critica", "pasta_analise_critica": o.obra,
                         "bairro": o.bairro, "unidades": o.unidades, "entrega_contratual": o.entrega.isoformat()})
        print(f"{o.obra:12} {o.bairro:12} início {o.inicio:%m/%Y} · término {o.termino:%m/%Y} · entrega "
              f"{o.entrega:%m/%Y} · exec set/26 {serie[-1].exec_acum:6.1%} · prev {serie[-1].prev_acum:6.1%} · "
              f"atraso {serie[-1].atraso:4.1f} m · IDC {serie[-1].valores['idc_alvo']:.3f}")
    crono = DADOS / "cronograma" / f"PLANILHA DE FLUXO FINANCEIRO ATUALIZADO EM {CORTE.month:02d}{CORTE.year}.xlsx"
    curvas = gravar_crono(por_obra, crono)
    gravar_csv(DADOS / "cadastro" / "obras_cadastro.csv", cadastro)
    gravar_csv(DADOS / "_gabarito" / "fato_analise_critica.csv", gab_ac)
    gravar_csv(DADOS / "_gabarito" / "base_complementar.csv", gab_base)
    gravar_csv(DADOS / "_gabarito" / "curva_prevista.csv", curvas)
    if not args.sem_recalculo:
        recalcular(arquivos)
    print(f"{len(arquivos) + len(OBRAS) + 1} planilhas · {len(gab_ac)} abas da Análise Crítica · {len(gab_base)} abas de base · "
          f"{len(curvas)} meses de curva prevista{' · fórmulas recalculadas (LibreOffice)' if not args.sem_recalculo else ''}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
